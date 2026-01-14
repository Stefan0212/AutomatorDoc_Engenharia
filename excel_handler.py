"""
Módulo para manipulação de arquivos Excel
"""
from openpyxl import load_workbook
from openpyxl.styles import Font
from io import BytesIO
import os
import zipfile
from dotenv import load_dotenv

load_dotenv()

# Senha do Excel (se necessário)
EXCEL_PASSWORD = os.getenv("EXCEL_PASSWORD", None)

# Linha padrão para preenchimento (linha abaixo dos cabeçalhos)
APPROVALS_DATA_ROW = 5


def has_vba_macros(file_content):
    """
    Verifica se o arquivo Excel contém macros VBA.
    Arquivos Excel são arquivos ZIP, então verificamos se contém arquivos VBA.
    
    Args:
        file_content (bytes): Conteúdo do arquivo Excel
    
    Returns:
        bool: True se o arquivo tem macros, False caso contrário
    """
    try:
        file_stream = BytesIO(file_content)
        with zipfile.ZipFile(file_stream, 'r') as zip_file:
            # Lista de arquivos/indicadores que sugerem a presença de macros VBA
            vba_indicators = [
                'xl/vbaProject.bin',  # Projeto VBA principal
                'xl/_rels/vbaProject.bin.rels',  # Relacionamentos do projeto VBA
            ]
            
            # Verifica se algum dos indicadores existe no arquivo ZIP
            for indicator in vba_indicators:
                if indicator in zip_file.namelist():
                    return True
            
            # Também verifica se há pastas/arquivos relacionados a VBA
            namelist = zip_file.namelist()
            for name in namelist:
                if 'vba' in name.lower() or 'macro' in name.lower():
                    return True
            
            return False
    except Exception:
        # Se não conseguir verificar (arquivo não é ZIP válido), assume que não tem macros
        return False


def read_excel_file(file_content):
    """
    Carrega um arquivo Excel a partir de conteúdo em bytes.
    Detecta automaticamente se o arquivo tem macros VBA e usa keep_vba apenas quando necessário.
    
    Args:
        file_content (bytes): Conteúdo do arquivo Excel
    
    Returns:
        openpyxl.Workbook: Workbook carregado
    """
    try:
        # Cria um novo BytesIO para evitar problemas de referência
        file_stream = BytesIO(file_content)
        
        # Detecta se o arquivo tem macros VBA
        has_macros = has_vba_macros(file_content)
        
        # Carrega o workbook com ou sem keep_vba dependendo da detecção
        if has_macros:
            # Arquivo tem macros: usa keep_vba=True para preservá-las
            workbook = load_workbook(
                filename=file_stream,
                data_only=False,  # Preserva fórmulas
                keep_vba=True  # Preserva macros VBA
            )
        else:
            # Arquivo não tem macros: carrega normalmente (mais seguro e rápido)
            workbook = load_workbook(
                filename=file_stream,
                data_only=False  # Preserva fórmulas
            )
        
        return workbook
    except Exception as e:
        raise Exception(f"Erro ao ler o arquivo Excel: {str(e)}")


def extract_document_info(workbook):
    """
    Extrai informações do documento da aba Cover usando busca dinâmica de rótulos.
    Percorre as primeiras 30 linhas e 10 colunas procurando pelos rótulos
    "DOCUMENT TITLE:" e "AIRCRAFT MODEL:" e extrai o valor da célula à direita.
    
    Args:
        workbook (openpyxl.Workbook): Workbook do Excel
    
    Returns:
        tuple: (document_title_key, aircraft_model)
    """
    try:
        if "Cover" not in workbook.sheetnames:
            raise ValueError("A aba 'Cover' não foi encontrada no arquivo Excel")
        
        cover_sheet = workbook["Cover"]
        
        # Dicionário para mapear rótulos encontrados
        labels_para_buscar = {
            "DOCUMENT TITLE:": "title",
            "AIRCRAFT MODEL:": "model"
        }
        
        metadata = {}
        
        # Busca dinâmica: percorre as primeiras 30 linhas e 10 colunas
        for row in cover_sheet.iter_rows(min_row=1, max_row=30, min_col=1, max_col=10):
            for cell in row:
                if cell.value:
                    cell_value = str(cell.value).strip().upper()
                    # Verifica se o valor da célula corresponde a algum rótulo
                    if cell_value in labels_para_buscar:
                        label_encontrado = cell_value
                        # O valor está na célula à direita (offset: col + 1)
                        valor_cell = cover_sheet.cell(row=cell.row, column=cell.column + 1)
                        valor = valor_cell.value
                        if valor:
                            # Normalização: strip e upper
                            valor_normalizado = str(valor).strip().upper()
                            metadata[labels_para_buscar[label_encontrado]] = valor_normalizado
        
        # Extrai os valores encontrados
        document_title_key = metadata.get("title", None)
        aircraft_model = metadata.get("model", None)
        
        # Validação: verifica se os rótulos foram encontrados
        if not document_title_key or not aircraft_model:
            raise ValueError("❌ Erro de Estrutura: Rótulos de identificação não localizados na aba Cover. Verifique se 'DOCUMENT TITLE:' e 'AIRCRAFT MODEL:' estão presentes.")
        
        return document_title_key, aircraft_model
    except ValueError:
        raise
    except Exception as e:
        raise Exception(f"Erro ao extrair informações do documento: {str(e)}")


def find_approval_headers(approvals_sheet):
    """
    Busca dinamicamente os cabeçalhos na aba Approvals.
    Procura pelos cabeçalhos: Prepared, Checked, Approved, Industrialization, Mat. Process
    
    Args:
        approvals_sheet: Worksheet da aba Approvals
    
    Returns:
        dict: Mapeamento {campo: coluna} onde cada campo corresponde à coluna do cabeçalho
    """
    # Mapeamento de cabeçalhos para campos internos
    headers_to_find = {
        "PREPARED": "prepared_by",
        "CHECKED": "checked_by",
        "APPROVED": "approved_by",
        "INDUSTRIALIZATION": "industrialization",
        "MAT. PROCESS": "materials_process",
        "MATERIALS PROCESS": "materials_process",  # Variação
        "MAT PROCESS": "materials_process"  # Variação
    }
    
    column_mapping = {}
    
    # Busca nas primeiras 10 linhas e 20 colunas (área típica de cabeçalhos)
    for row in approvals_sheet.iter_rows(min_row=1, max_row=10, min_col=1, max_col=20):
        for cell in row:
            if cell.value:
                cell_value = str(cell.value).strip().upper()
                # Verifica se o valor corresponde a algum cabeçalho
                for header_key, field_name in headers_to_find.items():
                    if header_key in cell_value:
                        # Encontrou o cabeçalho - salva a coluna
                        if field_name not in column_mapping:
                            column_mapping[field_name] = cell.column
                        break
    
    return column_mapping


def fill_approvals_sheet(workbook, approvers_data):
    """
    Preenche a aba Approvals com os dados dos aprovadores.
    Detecta automaticamente as colunas dos cabeçalhos e preenche os valores.
    
    Args:
        workbook (openpyxl.Workbook): Workbook do Excel
        approvers_data (dict): Dicionário com os dados dos aprovadores
    
    Returns:
        openpyxl.Workbook: Workbook atualizado
    """
    try:
        if "Approvals" not in workbook.sheetnames:
            raise ValueError("A aba 'Approvals' não foi encontrada no arquivo Excel")
        
        approvals_sheet = workbook["Approvals"]
        approvers = approvers_data.get("approvers", {})
        
        # Busca dinamicamente as colunas dos cabeçalhos
        column_mapping = find_approval_headers(approvals_sheet)
        
        # Verifica se encontrou todos os cabeçalhos necessários
        required_fields = ["prepared_by", "checked_by", "approved_by", "industrialization", "materials_process"]
        missing_headers = [field for field in required_fields if field not in column_mapping]
        
        if missing_headers:
            raise ValueError(
                f"❌ Erro: Não foi possível encontrar os seguintes cabeçalhos na aba Approvals: "
                f"{', '.join(missing_headers)}. Verifique se os cabeçalhos 'Prepared', 'Checked', "
                f"'Approved', 'Industrialization' e 'Mat. Process' estão presentes."
            )
        
        # Define a fonte padrão (Arial, tamanho 11)
        default_font = Font(name="Arial", size=11)
        
        # Preenche as células mapeadas dinamicamente
        for field_name, column in column_mapping.items():
            if field_name in approvers and approvers[field_name]:
                cell = approvals_sheet.cell(row=APPROVALS_DATA_ROW, column=column)
                cell.value = approvers[field_name]
                # Aplica a formatação
                cell.font = default_font
        
        return workbook
    except ValueError:
        raise
    except Exception as e:
        raise Exception(f"Erro ao preencher a aba Approvals: {str(e)}")


def save_workbook_to_bytes(workbook):
    """
    Salva o workbook em memória (BytesIO) para download.
    
    Args:
        workbook (openpyxl.Workbook): Workbook do Excel
    
    Returns:
        bytes: Conteúdo do arquivo Excel em bytes
    """
    try:
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output.getvalue()
    except Exception as e:
        raise Exception(f"Erro ao salvar o arquivo Excel: {str(e)}")


def process_excel_file(file_content, approvers_data):
    """
    Processa um arquivo Excel completo: lê, preenche e retorna os bytes.
    
    Args:
        file_content (bytes): Conteúdo do arquivo Excel original
        approvers_data (dict): Dicionário com os dados dos aprovadores
    
    Returns:
        bytes: Conteúdo do arquivo Excel processado
    """
    workbook = read_excel_file(file_content)
    workbook = fill_approvals_sheet(workbook, approvers_data)
    return save_workbook_to_bytes(workbook)
